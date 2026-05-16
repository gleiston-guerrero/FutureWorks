"""
process_metrics.py
==================
Lee los JSON generados por k6 (resultados_10vu.json, resultados_50vu.json,
resultados_100vu.json) y produce:
  - Tabla consolidada p50/p95/p99 con pandas
  - Excel de resultados para el paper
  - Gráfica de latencia por nivel de carga

Uso:
    python process_metrics.py               # busca JSONs en el directorio actual
    python process_metrics.py --dir ./runs  # directorio específico

Requerido por el spec del artículo: "procesar percentiles p50/p95/p99 con pandas"
"""

import json
import argparse
import os
import sys
from pathlib import Path

import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker

# ─────────────────────────────────────────────────────────────
# MAPEO DE NOMBRES (real → paper)
# ─────────────────────────────────────────────────────────────
ENDPOINT_ALIAS = {
    "POST /login/":                   "POST /login/",
    "POST /ml/recomendar-tipo-ruta":  "POST /rutas/generar",
    "GET /reminders/listar":          "GET /recordatorios",
    "WS connect time":                "WS /grupo/{id} — connect",
    "WS first message":               "WS /grupo/{id} — first msg",
}

CARGAS = [10, 50, 100]
SLA_P95 = {
    "POST /login/":                   60_000,
    "POST /ml/recomendar-tipo-ruta":   5_000,
    "GET /reminders/listar":           3_000,
    "WS connect time":                 5_000,
    "WS first message":                8_000,
}

# ─────────────────────────────────────────────────────────────
# CARGA DE DATOS
# ─────────────────────────────────────────────────────────────

def parse_ms(value: str) -> float:
    """Convierte '123 ms' o 'N/A' a float."""
    if value in ("N/A", None, ""):
        return float("nan")
    return float(str(value).replace(" ms", "").strip())


def load_results(directory: Path) -> dict:
    """
    Carga los 3 archivos JSON de k6 y devuelve un dict:
      {nivel_usuarios: {endpoint_name: {p50, p95, p99, avg, threshold}}}
    """
    data = {}
    for vu in CARGAS:
        path = directory / f"resultados_{vu}vu.json"
        if not path.exists():
            print(f"  [WARN] No encontrado: {path} — se omitirá esta carga")
            continue
        with open(path, encoding="utf-8") as f:
            raw = json.load(f)
        data[vu] = {}
        for ep_name, metrics in raw.get("endpoints", {}).items():
            data[vu][ep_name] = {
                "p50":       parse_ms(metrics.get("p50")),
                "p95":       parse_ms(metrics.get("p95")),
                "p99":       parse_ms(metrics.get("p99")),
                "avg":       parse_ms(metrics.get("avg")),
                "threshold": metrics.get("threshold", "N/A"),
            }
        data[vu]["_resumen"] = raw.get("resumen", {})
    return data


# ─────────────────────────────────────────────────────────────
# CONSTRUCCIÓN DE DATAFRAMES
# ─────────────────────────────────────────────────────────────

def build_latency_df(data: dict) -> pd.DataFrame:
    """
    DataFrame principal: una fila por (endpoint, nivel_carga).
    Columnas: endpoint_paper, endpoint_real, vu, p50, p95, p99, avg,
              sla_p95_ms, sla_ok, threshold_k6
    """
    rows = []
    for vu, endpoints in data.items():
        for ep_name, metrics in endpoints.items():
            if ep_name == "_resumen":
                continue
            rows.append({
                "endpoint_real":  ep_name,
                "endpoint_paper": ENDPOINT_ALIAS.get(ep_name, ep_name),
                "usuarios_virtuales": vu,
                "p50_ms":  metrics["p50"],
                "p95_ms":  metrics["p95"],
                "p99_ms":  metrics["p99"],
                "avg_ms":  metrics["avg"],
                "sla_p95_ms": SLA_P95.get(ep_name, float("nan")),
                "sla_ok":  (metrics["p95"] <= SLA_P95.get(ep_name, float("inf")))
                           if not pd.isna(metrics["p95"]) else None,
                "threshold_k6": metrics["threshold"],
            })
    df = pd.DataFrame(rows)
    df.sort_values(["endpoint_real", "usuarios_virtuales"], inplace=True)
    df.reset_index(drop=True, inplace=True)
    return df


def build_summary_df(data: dict) -> pd.DataFrame:
    """DataFrame de resumen global por nivel de carga."""
    rows = []
    for vu, endpoints in data.items():
        r = endpoints.get("_resumen", {})
        rows.append({
            "usuarios_virtuales":   vu,
            "http_requests":        r.get("total_peticiones_http", "N/A"),
            "ws_sessions":          r.get("total_conexiones_ws", "N/A"),
            "ws_errors":            r.get("errores_ws", "N/A"),
            "error_rate":           r.get("tasa_errores", "N/A"),
            "sla_error_rate":       r.get("threshold_error_rate", "N/A"),
        })
    df = pd.DataFrame(rows).sort_values("usuarios_virtuales").reset_index(drop=True)
    return df


def build_paper_table(df_latency: pd.DataFrame) -> pd.DataFrame:
    """
    Tabla en formato paper: filas = endpoint, columnas = (VU, p50/p95/p99).
    """
    endpoints_order = list(ENDPOINT_ALIAS.values())
    pivot = df_latency.pivot_table(
        index="endpoint_paper",
        columns="usuarios_virtuales",
        values=["p50_ms", "p95_ms", "p99_ms"],
        aggfunc="first",
    )
    # Reordenar columnas como (VU, métrica)
    cols = []
    for vu in CARGAS:
        for met in ["p50_ms", "p95_ms", "p99_ms"]:
            cols.append((met, vu))
    pivot = pivot.reindex(columns=pd.MultiIndex.from_tuples(cols))
    # Reordenar filas
    pivot = pivot.reindex([e for e in endpoints_order if e in pivot.index])
    pivot.columns = [f"{vu}VU_{m.replace('_ms','')}" for m, vu in pivot.columns]
    return pivot.round(0)


# ─────────────────────────────────────────────────────────────
# VISUALIZACIÓN
# ─────────────────────────────────────────────────────────────

def plot_latency(df_latency: pd.DataFrame, output_dir: Path):
    """
    Gráfica de barras agrupadas: p95 por endpoint y nivel de carga.
    """
    endpoints = list(ENDPOINT_ALIAS.values())
    df_plot = df_latency[df_latency["endpoint_paper"].isin(endpoints)].copy()

    fig, ax = plt.subplots(figsize=(12, 5))
    width   = 0.22
    x_base  = range(len(endpoints))
    colors  = {10: "#3498db", 50: "#e67e22", 100: "#e74c3c"}

    for i, vu in enumerate(CARGAS):
        subset = df_plot[df_plot["usuarios_virtuales"] == vu]
        vals   = []
        for ep in endpoints:
            row = subset[subset["endpoint_paper"] == ep]
            vals.append(row["p95_ms"].values[0] if not row.empty else float("nan"))
        positions = [x + (i - 1) * width for x in x_base]
        ax.bar(positions, vals, width=width, label=f"{vu} VUs",
               color=colors[vu], alpha=0.85, edgecolor="white")

    # SLA líneas
    for ep_real, sla in SLA_P95.items():
        ep_paper = ENDPOINT_ALIAS.get(ep_real, ep_real)
        if ep_paper in endpoints:
            xi = endpoints.index(ep_paper)
            ax.hlines(sla, xi - 1.5 * width, xi + 1.5 * width,
                      colors="black", linestyles="--", linewidth=0.8, alpha=0.5)

    ax.set_xticks(list(x_base))
    ax.set_xticklabels(endpoints, rotation=20, ha="right", fontsize=9)
    ax.set_ylabel("Latencia p95 (ms)", fontsize=11)
    ax.set_title("Latencia p95 por endpoint y nivel de carga\n"
                 "(línea punteada = SLA definido)", fontsize=12)
    ax.legend(fontsize=10)
    ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda x, _: f"{int(x):,}"))
    ax.grid(axis="y", alpha=0.3)
    plt.tight_layout()

    fig_path = output_dir / "latencia_p95_por_carga.png"
    plt.savefig(str(fig_path), dpi=150, bbox_inches="tight")
    plt.close()
    print(f"  ✓ Figura: {fig_path}")


# ─────────────────────────────────────────────────────────────
# EXCEL DE RESULTADOS
# ─────────────────────────────────────────────────────────────

def export_excel(df_latency: pd.DataFrame, df_summary: pd.DataFrame,
                 df_paper: pd.DataFrame, output_dir: Path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.utils.dataframe import dataframe_to_rows
    except ImportError:
        print("  [WARN] openpyxl no instalado — se omite Excel. pip install openpyxl")
        return

    wb = Workbook()
    hfont  = Font(bold=True, color="FFFFFF", size=11)
    hfill  = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )
    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font      = hfont
            cell.fill      = hfill
            cell.alignment = center
            cell.border    = border

    def style_data(ws, start_row=2):
        for row in ws.iter_rows(min_row=start_row):
            for cell in row:
                cell.alignment = center
                cell.border    = border

    def autowidth(ws):
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4

    # ── Hoja 1: Datos detallados ──────────────────────────
    ws1 = wb.active
    ws1.title = "Datos Detallados"
    cols_show = ["endpoint_paper", "usuarios_virtuales",
                 "p50_ms", "p95_ms", "p99_ms", "avg_ms",
                 "sla_p95_ms", "sla_ok", "threshold_k6"]
    headers1  = ["Endpoint (paper)", "VUs", "p50 (ms)", "p95 (ms)", "p99 (ms)",
                 "avg (ms)", "SLA p95 (ms)", "SLA OK", "Threshold k6"]

    for col, h in enumerate(headers1, 1):
        ws1.cell(row=1, column=col, value=h)
    style_header(ws1)

    for ri, row in enumerate(df_latency[cols_show].itertuples(index=False), 2):
        for ci, val in enumerate(row, 1):
            cell = ws1.cell(row=ri, column=ci, value=val)
        # Colorear SLA OK
        sla_cell = ws1.cell(row=ri, column=8)
        if sla_cell.value is True:
            sla_cell.fill = pass_fill
        elif sla_cell.value is False:
            sla_cell.fill = fail_fill

    style_data(ws1)
    autowidth(ws1)

    # ── Hoja 2: Tabla Paper ───────────────────────────────
    ws2 = wb.create_sheet("Tabla Paper")
    df_paper_reset = df_paper.reset_index()
    for col, h in enumerate(df_paper_reset.columns, 1):
        ws2.cell(row=1, column=col, value=h)
    style_header(ws2)
    for ri, row in enumerate(df_paper_reset.itertuples(index=False), 2):
        for ci, val in enumerate(row, 1):
            ws2.cell(row=ri, column=ci, value=val)
    style_data(ws2)
    autowidth(ws2)

    # ── Hoja 3: Resumen global ────────────────────────────
    ws3 = wb.create_sheet("Resumen Global")
    headers3 = ["VUs", "HTTP requests", "WS sessions", "WS errors",
                "Tasa errores", "SLA errores"]
    for col, h in enumerate(headers3, 1):
        ws3.cell(row=1, column=col, value=h)
    style_header(ws3)
    for ri, row in enumerate(df_summary.itertuples(index=False), 2):
        for ci, val in enumerate(row, 1):
            ws3.cell(row=ri, column=ci, value=val)
        # Colorear SLA
        sla_cell = ws3.cell(row=ri, column=6)
        if str(sla_cell.value).upper() == "PASS":
            sla_cell.fill = pass_fill
        elif str(sla_cell.value).upper() == "FAIL":
            sla_cell.fill = fail_fill
    style_data(ws3)
    autowidth(ws3)

    xlsx_path = output_dir / "latency_results.xlsx"
    wb.save(xlsx_path)
    print(f"  ✓ Excel: {xlsx_path}")


# ─────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Procesa resultados k6 con pandas")
    parser.add_argument("--dir", default=".", help="Directorio con los JSON de k6")
    args = parser.parse_args()

    input_dir  = Path(args.dir).resolve()
    output_dir = input_dir

    print(f"\nProcesando resultados en: {input_dir}")
    print("─" * 60)

    data = load_results(input_dir)
    if not data:
        print("ERROR: no se encontró ningún archivo resultados_*vu.json")
        sys.exit(1)

    df_latency = build_latency_df(data)
    df_summary = build_summary_df(data)
    df_paper   = build_paper_table(df_latency)

    # ── Consola: tabla paper ───────────────────────────────
    print("\n=== TABLA PAPER — Latencia por endpoint (ms) ===")
    print(df_paper.to_string())

    print("\n=== RESUMEN GLOBAL POR NIVEL DE CARGA ===")
    print(df_summary.to_string(index=False))

    print("\n=== CUMPLIMIENTO DE SLA (p95) ===")
    sla_check = df_latency[["endpoint_paper", "usuarios_virtuales", "p95_ms", "sla_p95_ms", "sla_ok"]].copy()
    sla_check["sla_ok"] = sla_check["sla_ok"].map({True: "✓ PASS", False: "✗ FAIL", None: "N/A"})
    print(sla_check.to_string(index=False))

    # ── CSV de salida ──────────────────────────────────────
    csv_path = output_dir / "latency_results.csv"
    df_latency.to_csv(csv_path, index=False, encoding="utf-8-sig")
    print(f"\n  ✓ CSV: {csv_path}")

    paper_csv = output_dir / "latency_paper_table.csv"
    df_paper.to_csv(paper_csv, encoding="utf-8-sig")
    print(f"  ✓ CSV tabla paper: {paper_csv}")

    # ── Figura ─────────────────────────────────────────────
    plot_latency(df_latency, output_dir)

    # ── Excel ─────────────────────────────────────────────
    export_excel(df_latency, df_summary, df_paper, output_dir)

    print("\n✅ Proceso completado.")


if __name__ == "__main__":
    main()